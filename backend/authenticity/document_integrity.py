import os
import re
import statistics
import zipfile
from collections import Counter
from datetime import datetime, timezone

import fitz
from docx import Document

from .utils import create_finding, safe_status


def _parse_pdf_date(value):
    if not value:
        return None
    match = re.search(r"(\d{4})(\d{2})(\d{2})(\d{2})?(\d{2})?(\d{2})?", str(value))
    if not match:
        return None
    parts = [int(part or 0) for part in match.groups()]
    try:
        return datetime(
            parts[0], parts[1], max(parts[2], 1),
            parts[3], parts[4], parts[5],
            tzinfo=timezone.utc,
        )
    except ValueError:
        return None


def _font_anomalies(font_sizes, font_names):
    findings = []
    if len(font_sizes) >= 10:
        minimum = min(font_sizes)
        maximum = max(font_sizes)
        median = statistics.median(font_sizes)
        if minimum > 0 and maximum / minimum >= 5 and maximum - minimum >= 12:
            findings.append(
                create_finding(
                    "document_integrity",
                    "Potentially Inconsistent",
                    "The PDF contains a large font-size spread that may warrant visual review.",
                    [
                        f"Observed font sizes range from {minimum:.1f} pt to {maximum:.1f} pt.",
                        f"Median observed font size is {median:.1f} pt.",
                    ],
                    0.68,
                    "Visually inspect the document for accidental or inserted formatting.",
                )
            )

    if len(font_names) >= 10:
        counts = Counter(font_names)
        rare_fonts = [name for name, count in counts.items() if count == 1]
        if rare_fonts and len(rare_fonts) >= max(2, int(len(counts) * 0.6)):
            findings.append(
                create_finding(
                    "document_integrity",
                    "Potentially Inconsistent",
                    "The PDF uses several one-off fonts; this is a formatting signal, not proof of tampering.",
                    [f"Rare fonts detected: {', '.join(rare_fonts[:8])}"],
                    0.52,
                    "Visually inspect unusual font changes around the affected text.",
                )
            )
    return findings


def _analyze_pdf(file_path):
    document = fitz.open(file_path)
    metadata = dict(document.metadata or {})
    page_count = document.page_count

    text_parts = []
    font_sizes = []
    font_names = []
    light_text_samples = []
    image_only_pages = 0

    for page_index, page in enumerate(document, start=1):
        page_text = page.get_text("text") or ""
        text_parts.append(page_text)

        if not page_text.strip() and page.get_images(full=True):
            image_only_pages += 1

        try:
            blocks = page.get_text("dict").get("blocks", [])
            for block in blocks:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        text = str(span.get("text") or "").strip()
                        if not text:
                            continue
                        size = float(span.get("size") or 0)
                        if size > 0:
                            font_sizes.append(size)
                        font = str(span.get("font") or "").strip()
                        if font:
                            font_names.append(font)

                        color = span.get("color")
                        if isinstance(color, int):
                            red = (color >> 16) & 255
                            green = (color >> 8) & 255
                            blue = color & 255
                            if min(red, green, blue) >= 245:
                                light_text_samples.append(
                                    f"Page {page_index}: '{text[:70]}'"
                                )
                        if size <= 2.5:
                            light_text_samples.append(
                                f"Page {page_index}: tiny text '{text[:70]}'"
                            )
        except Exception:
            continue

    findings = []

    if light_text_samples:
        findings.append(
            create_finding(
                "document_integrity",
                "Potentially Inconsistent",
                "Very light/white or extremely small text was detected and may be visually hidden in the document.",
                light_text_samples[:6],
                0.74,
                "Inspect the highlighted page regions manually; hidden-looking text can also be intentional design.",
            )
        )

    creation_date = _parse_pdf_date(metadata.get("creationDate"))
    modification_date = _parse_pdf_date(metadata.get("modDate"))
    now = datetime.now(timezone.utc)

    if creation_date and creation_date > now:
        findings.append(
            create_finding(
                "document_integrity",
                "Potentially Inconsistent",
                "The PDF creation timestamp is in the future relative to the analysis time.",
                [f"Creation timestamp: {creation_date.isoformat()}"],
                0.9,
                "Confirm the file timestamp and source document history.",
            )
        )

    if creation_date and modification_date and modification_date < creation_date:
        findings.append(
            create_finding(
                "document_integrity",
                "Potentially Inconsistent",
                "The PDF modification timestamp precedes its creation timestamp.",
                [
                    f"Creation timestamp: {creation_date.isoformat()}",
                    f"Modification timestamp: {modification_date.isoformat()}",
                ],
                0.86,
                "Confirm whether the source application rewrote PDF metadata.",
            )
        )

    embedded_names = []
    try:
        embedded_names = list(document.embfile_names())
    except Exception:
        embedded_names = []

    if embedded_names:
        findings.append(
            create_finding(
                "document_integrity",
                "Potentially Inconsistent",
                "The PDF contains embedded files.",
                [f"Embedded files: {', '.join(embedded_names[:10])}"],
                0.65,
                "Review embedded content and confirm it is expected for the resume.",
            )
        )

    try:
        with open(file_path, "rb") as handle:
            raw_bytes = handle.read()
        if b"/JavaScript" in raw_bytes or b"/JS" in raw_bytes:
            findings.append(
                create_finding(
                    "document_integrity",
                    "Potentially Inconsistent",
                    "PDF action/script markers were detected.",
                    ["PDF bytes contain JavaScript-related action markers."],
                    0.62,
                    "Open the PDF with a safe viewer and confirm no unexpected active content is present.",
                )
            )
    except OSError:
        pass

    findings.extend(_font_anomalies(font_sizes, font_names))

    status = "Potentially Inconsistent" if findings else ("Unverified" if image_only_pages else "Verified")
    note = (
        "No document-level anomaly signal was detected by the automated checks."
        if not findings
        else "Automated document checks found signals requiring human review."
    )

    document.close()

    return {
        "status": safe_status(status),
        "summary": note,
        "checks": {
            "file_type": "pdf",
            "page_count": page_count,
            "text_characters": sum(len(item) for item in text_parts),
            "image_only_pages": image_only_pages,
            "metadata": metadata,
            "font_count": len(set(font_names)),
            "font_sizes_sample": sorted({round(v, 1) for v in font_sizes})[:30],
            "embedded_files": embedded_names,
        },
        "findings": findings,
        "extracted_text": "\n".join(text_parts).strip(),
    }


def _analyze_docx(file_path):
    document = Document(file_path)
    core = document.core_properties
    findings = []
    text_parts = []

    for paragraph in document.paragraphs:
        text_parts.append(paragraph.text)
        for run in paragraph.runs:
            xml = run._r.xml
            if "<w:vanish" in xml:
                findings.append(
                    create_finding(
                        "document_integrity",
                        "Potentially Inconsistent",
                        "Hidden-run formatting was detected in the DOCX XML.",
                        [f"Hidden text run in paragraph: '{paragraph.text[:80]}'"],
                        0.77,
                        "Open the source document in Word and review hidden formatting.",
                    )
                )
                break

            color = run.font.color
            try:
                rgb = color.rgb if color else None
                if rgb:
                    hex_value = str(rgb)
                    if len(hex_value) == 6 and all(int(hex_value[i:i+2], 16) >= 245 for i in (0, 2, 4)):
                        if run.text.strip():
                            findings.append(
                                create_finding(
                                    "document_integrity",
                                    "Potentially Inconsistent",
                                    "A white/light font color was detected in a DOCX run.",
                                    [f"Light-colored text: '{run.text[:80]}'"],
                                    0.65,
                                    "Visually inspect the run; white text can be intentional in templates.",
                                )
                            )
            except Exception:
                pass

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)

    font_names = []
    font_sizes = []
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.name:
                font_names.append(run.font.name)
            if run.font.size:
                try:
                    font_sizes.append(float(run.font.size.pt))
                except Exception:
                    pass

    with zipfile.ZipFile(file_path) as archive:
        names = archive.namelist()
        embedded = [name for name in names if name.startswith("word/embeddings/")]
        has_vba = any("vbaProject.bin" in name for name in names)
        external_links = [name for name in names if name.startswith("word/externalLinks/")]

    if embedded:
        findings.append(
            create_finding(
                "document_integrity",
                "Potentially Inconsistent",
                "The DOCX contains embedded objects.",
                [f"Embedded object entries: {', '.join(embedded[:10])}"],
                0.62,
                "Confirm that embedded content is expected.",
            )
        )

    if external_links:
        findings.append(
            create_finding(
                "document_integrity",
                "Potentially Inconsistent",
                "The DOCX contains external-link relationships.",
                [f"External link entries: {len(external_links)}"],
                0.58,
                "Confirm that external links are intentional.",
            )
        )

    if has_vba:
        findings.append(
            create_finding(
                "document_integrity",
                "Potentially Inconsistent",
                "A VBA project was detected inside the document package.",
                ["word/vbaProject.bin is present."],
                0.92,
                "Do not open the document in an unsafe environment; confirm why macros are present.",
            )
        )

    if core.created and core.modified:
        created = core.created.astimezone(timezone.utc) if core.created.tzinfo else core.created.replace(tzinfo=timezone.utc)
        modified = core.modified.astimezone(timezone.utc) if core.modified.tzinfo else core.modified.replace(tzinfo=timezone.utc)
        if modified < created:
            findings.append(
                create_finding(
                    "document_integrity",
                    "Potentially Inconsistent",
                    "The DOCX modification time precedes its creation time.",
                    [f"Created: {created.isoformat()}", f"Modified: {modified.isoformat()}"],
                    0.86,
                    "Confirm the source application's metadata behavior.",
                )
            )

        if modified > datetime.now(timezone.utc):
            findings.append(
                create_finding(
                    "document_integrity",
                    "Potentially Inconsistent",
                    "The DOCX modification timestamp is in the future.",
                    [f"Modified: {modified.isoformat()}"],
                    0.9,
                    "Confirm the source file timestamp.",
                )
            )

    if font_sizes:
        findings.extend(_font_anomalies(font_sizes, font_names))

    text = "\n".join(part for part in text_parts if part.strip()).strip()
    status = "Potentially Inconsistent" if findings else "Verified"

    return {
        "status": safe_status(status),
        "summary": (
            "No document-level anomaly signal was detected by the automated checks."
            if not findings
            else "Automated document checks found signals requiring human review."
        ),
        "checks": {
            "file_type": "docx",
            "text_characters": len(text),
            "metadata": {
                "title": core.title,
                "subject": core.subject,
                "author": core.author,
                "last_modified_by": core.last_modified_by,
                "created": core.created.isoformat() if core.created else None,
                "modified": core.modified.isoformat() if core.modified else None,
            },
            "font_count": len(set(font_names)),
            "font_sizes_sample": sorted({round(v, 1) for v in font_sizes})[:30],
            "embedded_objects": embedded,
            "external_links": external_links[:20],
            "has_vba": has_vba,
        },
        "findings": findings,
        "extracted_text": text,
    }


def analyze_document(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".pdf":
        return _analyze_pdf(file_path)
    if extension == ".docx":
        return _analyze_docx(file_path)
    raise ValueError("Unsupported resume format. Only PDF and DOCX are supported.")
